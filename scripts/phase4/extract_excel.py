#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Phase 4 管线基础设施 · Excel 抽取器
====================================

职责（只做原始落盘，不做任何科目语义转换）：
  1. 用 openpyxl 读取《formula-challenge-V1V2数据总览.xlsx》的四张主表 + 教材原文参照表；
  2. 单元格字符串前后空白一律剥离；剥离后为空串 / None 一律规范化为 null；
  3. 每条记录保留 sourceSheet（源 sheet 名）与 sourceRow（源 Excel 行号，1-based，
     含表头偏移），供后续质量报告精确定位；
  4. 输出到 data/raw/*.json，供 normalize 阶段（TS 共享库）消费。

输出文件（与 PIPELINE-SPEC.md 对齐）：
  data/raw/formula_full.json           V1_方剂完整（190 行）
  data/raw/formula_pending.json        V1_方剂待补全（70 行）
  data/raw/formula_textbook_ref.json   V1_方剂教材原文（212 行，仅交叉去重参照，不入主导入）
  data/raw/herbs.json                  V1_中药（306 行）
  data/raw/acupoints.json              V1_腧穴（259 行）

运行：python scripts/phase4/extract_excel.py
"""

import json
import sys
from datetime import datetime, timezone
from pathlib import Path

import openpyxl

# 项目根 = scripts/phase4/ 的上两级
ROOT = Path(__file__).resolve().parents[2]
XLSX = ROOT.parent / "formula-challenge-V1V2数据总览.xlsx"
RAW_DIR = ROOT / "data" / "raw"


def norm(v):
    """规范化单元格值：剥空白；空白串/None → None；其余原样保留（int 等不动）。"""
    if v is None:
        return None
    if isinstance(v, str):
        s = v.strip()
        return s if s else None
    return v


def extract_sheet(ws, field_map):
    """按 field_map {中文字段名: 0-based 列号} 抽取，跳过全空行。"""
    records = []
    for row_index, row in enumerate(
        ws.iter_rows(min_row=2, values_only=True), start=2
    ):
        if not any(v is not None and str(v).strip() for v in row):
            continue
        fields = {name: norm(row[col]) for name, col in field_map.items()}
        records.append(
            {
                "sourceSheet": ws.title,
                "sourceRow": row_index,
                "fields": fields,
            }
        )
    return records


# 各表列映射（0-based，已用 openpyxl 核实表头）
FORMULA_FULL_COLS = {
    "名称": 0, "章": 1, "章节名": 2, "等级": 3, "组成": 4,
    "功效": 5, "主治": 6, "助记": 7, "助记解释": 8,
    "传统方歌": 9, "传统方歌解释": 10, "触发词": 11, "数据来源": 12,
}
FORMULA_PENDING_COLS = {
    "名称": 0, "章": 1, "章节名": 2, "等级": 3, "组成": 4,
    "功效": 5, "主治": 6, "方歌": 7,
}
FORMULA_TEXTBOOK_REF_COLS = {
    "名称": 0, "组成": 1, "功效": 2, "主治": 3, "方歌": 4,
}
HERBS_COLS = {
    "ID": 0, "名称": 1, "分类": 2, "性味": 3, "归经": 4,
    "功效": 5, "主治": 6, "等级": 7, "助记": 8, "助记解释": 9, "数据来源": 10,
}
ACUPOINTS_COLS = {
    "ID": 0, "名称": 1, "拼音": 2, "编码": 3, "经络": 4,
    "定位": 5, "主治": 6, "刺法": 7, "特殊": 8, "注意事项": 9,
    "助记": 10, "助记解释": 11, "等级": 12, "排序": 13,
}


def main():
    if not XLSX.exists():
        sys.exit(f"[FATAL] 找不到 Excel 源文件：{XLSX}")
    RAW_DIR.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    generated_at = datetime.now(timezone.utc).isoformat()

    jobs = [
        ("formula_full", "V1_方剂完整", FORMULA_FULL_COLS),
        ("formula_pending", "V1_方剂待补全", FORMULA_PENDING_COLS),
        ("formula_textbook_ref", "V1_方剂教材原文", FORMULA_TEXTBOOK_REF_COLS),
        ("herbs", "V1_中药", HERBS_COLS),
        ("acupoints", "V1_腧穴", ACUPOINTS_COLS),
    ]

    summary = {}
    for kind, sheet_name, cols in jobs:
        if sheet_name not in wb.sheetnames:
            sys.exit(f"[FATAL] sheet 不存在：{sheet_name}")
        records = extract_sheet(wb[sheet_name], cols)
        payload = {
            "kind": kind,
            "sourceFile": XLSX.name,
            "sourceSheet": sheet_name,
            "generatedAt": generated_at,
            "count": len(records),
            "records": records,
        }
        out_path = RAW_DIR / f"{kind}.json"
        with out_path.open("w", encoding="utf-8") as f:
            json.dump(payload, f, ensure_ascii=False, indent=2)
        summary[kind] = {"file": str(out_path.name), "count": len(records)}
        print(f"[OK] {sheet_name:14s} -> {out_path.name:28s} {len(records)} 条")

    # 打印首条样例供人工核验
    print("\n===== 首条样例核验 =====")
    for kind in summary:
        with (RAW_DIR / f"{kind}.json").open("r", encoding="utf-8") as f:
            payload = json.load(f)
        first = payload["records"][0] if payload["records"] else None
        print(f"\n--- {kind} (count={payload['count']}) ---")
        print(json.dumps(first, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    # Windows 控制台中文输出
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass
    main()
